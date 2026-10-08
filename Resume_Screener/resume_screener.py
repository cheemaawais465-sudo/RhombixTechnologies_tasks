"""
Automated Resume Screener
-------------------------
Scans multiple resumes (PDF / DOCX), filters them against job requirements
(keywords, experience, education), scores each candidate and generates a
report of shortlisted candidates (Excel + text). Optional email alert.

Usage:
    python resume_screener.py --resumes ./resumes --config job_requirements.json
    python resume_screener.py --resumes ./resumes --config job_requirements.json --email hr@company.com
"""

import argparse
import json
import os
import re
import smtplib
import sys
from datetime import datetime
from email.message import EmailMessage

import pandas as pd

try:
    import pdfplumber
except ImportError:
    pdfplumber = None
try:
    import docx  # python-docx
except ImportError:
    docx = None


# ----------------------------------------------------------------------------
# 1. Text extraction
# ----------------------------------------------------------------------------
def extract_text_pdf(path: str) -> str:
    if pdfplumber is None:
        raise RuntimeError("pdfplumber not installed: pip install pdfplumber")
    text = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text.append(page.extract_text() or "")
    return "\n".join(text)


def extract_text_docx(path: str) -> str:
    if docx is None:
        raise RuntimeError("python-docx not installed: pip install python-docx")
    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:  # many CVs use tables for layout
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def extract_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return extract_text_pdf(path)
    if ext == ".docx":
        return extract_text_docx(path)
    raise ValueError(f"Unsupported file type: {ext}")


# ----------------------------------------------------------------------------
# 2. Information extraction
# ----------------------------------------------------------------------------
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[\s-]?)?(?:\(?\d{3,4}\)?[\s-]?)\d{3}[\s-]?\d{4}")
EXP_RE = re.compile(r"(\d{1,2})\s*\+?\s*(?:years?|yrs?)", re.IGNORECASE)


def find_email(text: str) -> str:
    m = EMAIL_RE.search(text)
    return m.group(0) if m else ""


def find_phone(text: str) -> str:
    m = PHONE_RE.search(text)
    return m.group(0).strip() if m else ""


def guess_name(text: str, filename: str) -> str:
    # First non-empty line that looks like a name, else fall back to filename
    for line in text.splitlines():
        line = line.strip()
        if 2 <= len(line.split()) <= 4 and re.fullmatch(r"[A-Za-z .'-]+", line):
            return line.title()
    return os.path.splitext(filename)[0].replace("_", " ").title()


def find_experience_years(text: str) -> int:
    years = [int(y) for y in EXP_RE.findall(text)]
    years = [y for y in years if y <= 40]  # ignore silly numbers
    return max(years) if years else 0


def keyword_present(keyword: str, text: str) -> bool:
    pattern = r"(?<!\w)" + re.escape(keyword.lower()) + r"(?!\w)"
    return re.search(pattern, text.lower()) is not None


# ----------------------------------------------------------------------------
# 3. Scoring
# ----------------------------------------------------------------------------
def score_resume(text: str, cfg: dict) -> dict:
    required = cfg.get("required_keywords", [])
    optional = cfg.get("optional_keywords", [])
    education = cfg.get("education_keywords", [])
    min_exp = cfg.get("min_experience_years", 0)
    w = cfg.get("weights", {"required": 60, "optional": 20, "education": 10, "experience": 10})

    req_found = [k for k in required if keyword_present(k, text)]
    req_missing = [k for k in required if k not in req_found]
    opt_found = [k for k in optional if keyword_present(k, text)]
    edu_found = [k for k in education if keyword_present(k, text)]
    exp_years = find_experience_years(text)

    req_score = w["required"] * (len(req_found) / len(required)) if required else w["required"]
    opt_score = w["optional"] * (len(opt_found) / len(optional)) if optional else 0
    edu_score = w["education"] if (edu_found or not education) else 0
    if min_exp <= 0:
        exp_score = w["experience"]
    else:
        exp_score = w["experience"] * min(exp_years / min_exp, 1)

    total = round(req_score + opt_score + edu_score + exp_score, 1)
    return {
        "score": total,
        "matched_required": ", ".join(req_found),
        "missing_required": ", ".join(req_missing),
        "matched_optional": ", ".join(opt_found),
        "education_match": ", ".join(edu_found),
        "experience_years": exp_years,
    }


# ----------------------------------------------------------------------------
# 4. Reports
# ----------------------------------------------------------------------------
def save_excel(df_all: pd.DataFrame, df_short: pd.DataFrame, path: str) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df_all.to_excel(writer, sheet_name="All Candidates", index=False)
        df_short.to_excel(writer, sheet_name="Shortlisted", index=False)
        for sheet in writer.sheets.values():
            for col in sheet.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                sheet.column_dimensions[col[0].column_letter].width = min(width + 2, 50)


def build_text_report(df_all: pd.DataFrame, df_short: pd.DataFrame, cfg: dict) -> str:
    lines = [
        "=" * 60,
        f"RESUME SCREENING REPORT - {cfg.get('job_title', 'Job')}",
        f"Generated: {datetime.now():%Y-%m-%d %H:%M}",
        "=" * 60,
        f"Total resumes scanned : {len(df_all)}",
        f"Shortlisted           : {len(df_short)} (threshold: {cfg.get('shortlist_threshold', 60)})",
        "-" * 60,
    ]
    if df_short.empty:
        lines.append("No candidate met the shortlist criteria.")
    for i, row in enumerate(df_short.itertuples(index=False), 1):
        lines += [
            f"{i}. {row.name}  |  Score: {row.score}",
            f"   Email: {row.email}  Phone: {row.phone}",
            f"   Matched: {row.matched_required}",
            f"   Missing: {row.missing_required or '-'}",
            f"   Experience: {row.experience_years} yrs  File: {row.file}",
            "",
        ]
    return "\n".join(lines)


def send_email_alert(to_addr: str, subject: str, body: str, attachment: str = None) -> None:
    """Needs env vars: SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASS (use Gmail app password)."""
    host = os.getenv("SMTP_HOST", "smtp.gmail.com")
    port = int(os.getenv("SMTP_PORT", "587"))
    user, pwd = os.getenv("SMTP_USER"), os.getenv("SMTP_PASS")
    if not user or not pwd:
        print("[!] SMTP_USER / SMTP_PASS not set - email skipped.")
        return
    msg = EmailMessage()
    msg["From"], msg["To"], msg["Subject"] = user, to_addr, subject
    msg.set_content(body)
    if attachment and os.path.exists(attachment):
        with open(attachment, "rb") as f:
            msg.add_attachment(
                f.read(), maintype="application", subtype="octet-stream",
                filename=os.path.basename(attachment),
            )
    with smtplib.SMTP(host, port) as server:
        server.starttls()
        server.login(user, pwd)
        server.send_message(msg)
    print(f"[+] Email alert sent to {to_addr}")


# ----------------------------------------------------------------------------
# 5. Main
# ----------------------------------------------------------------------------
def main() -> None:
    parser = argparse.ArgumentParser(description="Automated Resume Screener")
    parser.add_argument("--resumes", required=True, help="Folder containing PDF/DOCX resumes")
    parser.add_argument("--config", required=True, help="Job requirements JSON file")
    parser.add_argument("--out", default="output", help="Output folder")
    parser.add_argument("--email", help="Send shortlist report to this email")
    args = parser.parse_args()

    with open(args.config, encoding="utf-8") as f:
        cfg = json.load(f)
    threshold = cfg.get("shortlist_threshold", 60)

    files = [f for f in os.listdir(args.resumes) if f.lower().endswith((".pdf", ".docx"))]
    if not files:
        sys.exit("No PDF/DOCX resumes found in the folder.")

    rows = []
    for fname in sorted(files):
        path = os.path.join(args.resumes, fname)
        try:
            text = extract_text(path)
            if not text.strip():
                print(f"[!] {fname}: no text found (scanned image?) - skipped")
                continue
            result = score_resume(text, cfg)
            rows.append({
                "name": guess_name(text, fname),
                "email": find_email(text),
                "phone": find_phone(text),
                **result,
                "file": fname,
            })
            print(f"[+] {fname}: score {result['score']}")
        except Exception as e:  # keep going if one file is broken
            print(f"[!] {fname}: {e}")

    if not rows:
        sys.exit("No resumes could be processed.")

    df_all = pd.DataFrame(rows).sort_values("score", ascending=False).reset_index(drop=True)
    df_short = df_all[
        (df_all["score"] >= threshold) & (df_all["missing_required"] == "")
        | (df_all["score"] >= threshold + 15)
    ].reset_index(drop=True)

    os.makedirs(args.out, exist_ok=True)
    excel_path = os.path.join(args.out, "shortlisted_candidates.xlsx")
    txt_path = os.path.join(args.out, "report.txt")
    save_excel(df_all, df_short, excel_path)
    report = build_text_report(df_all, df_short, cfg)
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(report)

    print("\n" + report)
    print(f"[+] Excel saved: {excel_path}\n[+] Report saved: {txt_path}")

    if args.email:
        send_email_alert(
            args.email,
            f"Shortlisted candidates - {cfg.get('job_title', 'Job')}",
            report,
            excel_path,
        )


if __name__ == "__main__":
    main()
